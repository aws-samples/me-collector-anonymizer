#Beginning of Script
#!/usr/bin/env python3
import csv
import zipfile
import argparse
import openpyxl


def get_column_indexes(sheets):
    """Create a dict containing sheet name -> column name -> index # for all sheets provided."""
    headers = {}
    for sheet in sheets:
        headers[sheet.title] = {}
        for idx, column in enumerate(sheet.columns):
            headers[sheet.title][column[0].value] = idx + 1
    return headers


def anonymize(filename):
    """Anonymize the given Excel file, saving output as a new file."""
    wb = openpyxl.load_workbook(filename)

    # Read sheet names
    uti = wb["Utilization"]
    asset = wb["Asset Ownership"]
    virt = wb["Virtual Provisioning"]
    phys = wb["Physical Provisioning"]

    column_indexes = get_column_indexes([uti, asset, virt, phys])

    # replace Hypervisor Name in Virtual Provisioning sheet with hypervisor's Unique Identifier
    for cell in list(virt.columns)[
        column_indexes["Virtual Provisioning"]["Hypervisor Name"] - 1
    ]:
        for b_cell in list(phys.columns)[
            column_indexes["Physical Provisioning"]["Human Name"] - 1
        ]:
            if b_cell.value == cell.value:
                virt.cell(
                    row=cell.row,
                    column=column_indexes["Virtual Provisioning"]["Hypervisor Name"],
                ).value = phys.cell(
                    row=b_cell.row,
                    column=column_indexes["Physical Provisioning"]["Unique Identifier"],
                ).value

    # replace "Human Name" with "Unique Identifier" across all sheets
    for sheet in [uti, asset, virt, phys]:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(
                row=row, column=column_indexes[sheet.title]["Human Name"]
            ).value = sheet.cell(
                row=row, column=column_indexes[sheet.title]["Unique Identifier"]
            ).value

    # Remove IPs in Physical, Virtual Provisioning sheets
    for sheet in [phys, virt]:
        for cell in list(sheet.columns)[column_indexes[sheet.title]["Address"] - 1][1:]:
            cell.value = None

    wb.save("Inventory_And_Usage_Workbook Anonymized.xlsx")
    print(
        "Anonymization successful, Inventory_And_Usage_Workbook Anonymized has been created"
    )


def deanonymize(filename, qi):
    """De-anonymize the given Quick Insights .zip file using the original collector export file."""
    # Get filenames from the Quick Insights zip file
    with zipfile.ZipFile(qi, "r") as z:
        zip_filenames = z.namelist()

        # Get the Asset Ownership and Physical Provisioning of the original pre-anonymized script
        wb = openpyxl.load_workbook(filename)
        asset = wb["Asset Ownership"]
        phys = wb["Physical Provisioning"]

        column_indexes = get_column_indexes([asset, phys])

        unique_id_to_hostname = {}
        for row in asset.rows:
            unique_id_to_hostname[
                row[
                    column_indexes["Asset Ownership"]["Unique Identifier"] - 1
                ].value.upper()
            ] = row[column_indexes["Asset Ownership"]["Human Name"] - 1].value
        for row in phys.rows:
            unique_id_to_hostname[
                row[
                    column_indexes["Physical Provisioning"]["Unique Identifier"] - 1
                ].value.upper()
            ] = row[column_indexes["Physical Provisioning"]["Human Name"] - 1].value

        # create de-anonymized Server and SQL QI files
        for zip_file in zip_filenames:
            with z.open(zip_file) as f:
                file_string = f.read().decode("utf-8")

            reader = csv.reader(file_string.splitlines())
            headers = next(reader)
            csv_col_indexes = {}
            for idx, column in enumerate(headers):
                csv_col_indexes[column] = idx

            # Creating the de-anon file
            output_filename = "deanonymized_" + zip_file
            print(f"Creating output file: {output_filename}")

            with open(output_filename, "w", newline="") as g:
                writer = csv.writer(g)
                writer.writerow(headers)

                for row in reader:
                    # swap the "Server Name" column values back to the original hostnames
                    row[csv_col_indexes["Server Name"]] = unique_id_to_hostname[
                        row[csv_col_indexes["Server Id"]].upper()
                    ]

                    # swap the "Virtualization | Host Name" column values back to the original hostnames (if the column exists and has a value)
                    if (
                        "Virtualization | Host Name" in csv_col_indexes
                        and row[csv_col_indexes["Virtualization | Host Name"]]
                    ):
                        row[
                            csv_col_indexes["Virtualization | Host Name"]
                        ] = unique_id_to_hostname[
                            row[csv_col_indexes["Virtualization | Host Name"]].upper()
                        ]

                    # write modified row to CSV
                    writer.writerow(row)


if __name__ == "__main__":
    # arguments processing
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "method",
        help="The method to use, 'an' for anonymization or 'de' for de-anonymization",
    )
    parser.add_argument("filename", help="Inventory and Utilization Export file")
    parser.add_argument("QI", help="QI .zip file", nargs="?", default=None)
    args = parser.parse_args()

    if args.method == "an":
        anonymize(args.filename)

    elif args.method == "de":
        if args.QI is None:
            parser.error("QI is required for de-anonymization")
        deanonymize(args.filename, args.QI)

    else:
        parser.error("Invalid input. Please enter either 'an' or 'de'.")
#End of script
